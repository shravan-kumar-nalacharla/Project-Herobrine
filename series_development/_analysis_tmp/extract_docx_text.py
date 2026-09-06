from pathlib import Path
from docx import Document

source = Path(r"G:\herobrne smp script\Minecraft_The_Knot_Series_Bible.docx")
output = Path(r"G:\herobrne smp script\series_development\_analysis_tmp\old_bible_text.txt")
doc = Document(source)
lines = []
for p in doc.paragraphs:
    text = p.text.strip()
    if text:
        lines.append(text)
for index, table in enumerate(doc.tables, start=1):
    lines.append(f"\n[TABLE {index}]")
    for row in table.rows:
        lines.append(" | ".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
output.write_text("\n".join(lines), encoding="utf-8")
print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} lines={len(lines)} output={output}")
